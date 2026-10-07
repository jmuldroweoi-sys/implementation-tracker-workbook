"""Build workbook/implementation-tracker-workbook.xlsx reproducibly from pinned R1 inputs.

Steps: load the pinned R1 contracts and synthetic data, load R3 configuration, create the
12 sheets in order, create the Excel Tables, write inputs and formulas, apply styles,
formats, data validation, named ranges, frozen panes, filters, widths, protection, and
metadata, then save with fixed timestamps so the same inputs give the same bytes.

Usage:
    python tools/build_workbook.py                  build the workbook
    python tools/build_workbook.py --exports        build, recalculate a copy with
                                                    LibreOffice, and write exports/
    python tools/build_workbook.py --output PATH [--data-dir DIR] [--capacity CSV] [--empty]
                                                    build a variant (used by tests)

No macros, no external links, no volatile functions.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

# openpyxl serializes XML differently when the optional lxml package is installed.
# The builder always uses openpyxl's standard writer, so the bytes do not depend on
# whether lxml happens to be present. build() refuses to run if this did not take effect.
os.environ["OPENPYXL_LXML"] = "False"

import openpyxl.xml as openpyxl_xml  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.chart import BarChart, Reference  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill, Protection  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.workbook.defined_name import DefinedName  # noqa: E402
from openpyxl.workbook.properties import CalcProperties  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402
from openpyxl.worksheet.table import Table, TableColumn, TableFormula, TableStyleInfo  # noqa: E402

import reference_model as RM  # noqa: E402
import workbook_spec as W  # noqa: E402

HEADER_ROW = 4
FIRST_ROW = HEADER_ROW + 1
VALIDATION_ROWS = 1000
FONT = "Arial"
FILL_INPUT_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_FORMULA_HEAD = PatternFill("solid", fgColor="595959")
FILL_MIRROR_HEAD = PatternFill("solid", fgColor="375623")
FILL_FORMULA = PatternFill("solid", fgColor="EDEDED")
FILL_MIRROR = PatternFill("solid", fgColor="EEF3E8")
FILL_WARN = PatternFill("solid", fgColor="FCE4D6")
FILL_BAD = PatternFill("solid", fgColor="F8CBAD")
WHITE_BOLD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
BODY = Font(name=FONT, size=10)
FORMULA_FONT = Font(name=FONT, size=10, color="404040")
TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
SUBTITLE = Font(name=FONT, size=11, bold=True, color="1F3864")
NOTE = Font(name=FONT, size=10, italic=True, color="404040")

FIELD_VALIDATION = {
    "project_status": "project_status", "org_stage": "org_stage", "lifecycle_phase_key": "phase_key", "gate_id": "gate_id",
    "outcome": "gate_outcome", "request_type": "request_type", "priority": "request_priority", "source": "request_source",
    "handoff_type": "handoff_type", "severity": "severity", "category": "readiness_category", "assessment_purpose": "assessment_purpose",
    "complexity_profile_id": "profile_id", "service_profile_id": "profile_id", "segment_profile_id": "profile_id",
}
STATUS_VALIDATION = {
    "tblPhases": "phase_status", "tblMilestones": "milestone_status", "tblTasks": "task_status", "tblRequests": "request_status",
    "tblHandoffs": "handoff_status", "tblRisks": "risk_status", "tblIssues": "issue_status",
}
SHEET_NOTES = {
    "Lookups": "Every list is synchronized from the pinned R1 version (source columns show the file). Read-only; change R1, then re-sync and rebuild.",
    "Projects": "Grey columns are formulas (R3 checks). White columns are R1 project fields you edit.",
    "Phases and Gates": "A person records every gate outcome in tblGateAssessments. The workbook checks values; it never decides a gate.",
    "Tasks": "due_state uses calculation_as_of_at on the README tab, never the clock.",
    "Requests and Handoffs": "Statuses and legal transitions come from R1 config/request-state-machine.yaml. Escalation times come from R1 config/sla-rules.yaml.",
    "Escalation Rules": "Source of truth: R1 configuration. This tab is a read-only mirror; editing it would not change R1.",
    "Risks and Issues": "risk_score is calculated (likelihood x impact on the R1 scale); never type it. Severity values come from the R1 standard.",
    "Readiness Scorecard": "Readiness score informs human review. It does not decide go or no-go.",
    "Capacity Inputs": "Prepares workload hours for future capacity planning. It does not calculate capacity, utilization, staffing need, hiring timing, or capacity gap.",
    "Event Log": "Shared event contract (10 columns). Rows are imported or appended by a person; nothing is appended automatically. Check columns flag problems.",
    "KPI Summary": "Operational view traced to R1 records. Not a metric authority, not a forecast, and not a performance measure.",
}
SHEET_TITLES = {name: name for name in W.SHEETS}
SHEET_TITLES["README"] = "Implementation Tracker Workbook"


def readme_lines(st: dict) -> list[tuple[str, str]]:
    ref = W.pin_reference()
    return [
        ("h", "What is this workbook?"),
        ("p", "A macro-free tracker that applies the R1 implementation operating model (implementation-operating-system) to day-to-day project work. It is a reference implementation of R1, not a separate method."),
        ("p", "The bundled rows are synthetic data: three synthetic projects from R1 (Synthetic Projects A, B, and C). They are an illustrative example of the model in use, not customer data and not a record of historical use."),
        ("h", "What does R1 own?"),
        ("p", "The lifecycle and phase statuses, gate outcomes, project, task, request, handoff, risk, and issue statuses, request transitions, severity definitions, the risk formula, readiness categories and weights, SLA and lead-time parameters, and the event catalog."),
        ("h", "What does R3 (this workbook) own?"),
        ("p", "Workbook structure, the formulas that implement R1 rules, check columns, imports and exports, workbook metadata, and the Capacity Inputs export format."),
        ("h", "Which R1 version is pinned?"),
        ("p", f"R1 commit {ref['r1_commit']} (repository version {ref['r1_repository_version']}), shared standard {ref['shared_standard_version']}. See standard/standard-reference.yaml."),
        ("h", "What organization stage is selected?"),
        ("p", "The value selected_org_stage in the settings table. It sets which fields and controls config/stage-requirements.yaml treats as required, recommended, or optional. It never changes IDs, the event contract, readiness categories, or human gate authority."),
        ("h", "What is calculation_as_of_at?"),
        ("p", "The evaluation time for every due-state and escalation formula, in ISO 8601 UTC (for example 2026-10-05T23:59:59Z). It is a user-configurable parameter. No formula reads the computer clock, so the same inputs always give the same results."),
        ("h", "How do I use Starter Mode?"),
        ("p", "Set operating_mode to starter. Each week update Projects, Tasks, Risks and Issues, and Requests and Handoffs; review Readiness when a launch approaches; add Event Log rows only for material changes; then read KPI Summary. See docs/starter-mode.md."),
        ("h", "How do I use Full Mode?"),
        ("p", "Set operating_mode to full. Use all 12 tabs: complete gate evidence, capacity exports, and full event reconciliation. Same workbook, same IDs, same R1 rules. See docs/workbook-guide.md."),
        ("h", "Which cells do I edit?"),
        ("p", "White cells with dark blue headers: record fields on the record tabs, and the value column of the settings table below. Use the drop-down lists; they come from R1."),
        ("h", "Which cells contain formulas?"),
        ("p", "Grey cells with grey headers. They reproduce R1 rules or check your records. Do not type over them. Green-headed tables are read-only mirrors of R1 configuration."),
        ("h", "How do I export data?"),
        ("p", "Run python tools/build_workbook.py --exports, or save each table as CSV using the column list in config/export-map.yaml. Helper columns never enter authoritative exports. See docs/export-model.md."),
        ("h", "How do I maintain the Event Log?"),
        ("p", "After a material change (a new project, a phase move, a gate outcome, a status change, a new risk or issue), append one Event Log row using config/event-mappings.yaml, or import rows. The event_evidence columns show which changes still need a row. See docs/event-capture-model.md."),
        ("h", "What does the workbook intentionally not automate?"),
        ("p", "It does not append events, decide gates or launches, change statuses, read the clock, calculate capacity or staffing, or use AI. A named person makes every decision."),
    ]


def style_header(cell, kind: str) -> None:
    cell.font = WHITE_BOLD
    cell.fill = {"input": FILL_INPUT_HEAD, "formula": FILL_FORMULA_HEAD, "mirror": FILL_MIRROR_HEAD}[kind]
    cell.alignment = Alignment(wrap_text=True, vertical="center")


def cell_value(v):
    if isinstance(v, list):
        return ";".join(v)
    return v


def place_tables(sheet_specs: list[W.TableSpec], start_col: int = 1) -> dict[str, int]:
    positions, col = {}, start_col
    for spec in sheet_specs:
        positions[spec.name] = col
        col += len(spec.columns) + 1
    return positions


# Free-text columns that hold sentences get room to show them without wrapping.
WIDE_TEXT_COLUMNS = {"name": 40, "title": 34, "description": 44}


def display_width(column) -> float:
    """Column width that never splits a header name mid-word and gives free text room.
    Header names are capped at 34 characters of width; longer ones wrap at underscores."""
    return max(column.width, min(len(column.name) + 3, 34), WIDE_TEXT_COLUMNS.get(column.name, 0))


def write_table(ws, spec: W.TableSpec, rows: list[dict], first_col: int, wb, defined: list) -> tuple[int, int]:
    ws.cell(row=HEADER_ROW - 1, column=first_col, value=spec.title).font = SUBTITLE
    for j, column in enumerate(spec.columns):
        c = ws.cell(row=HEADER_ROW, column=first_col + j, value=column.name)
        style_header(c, column.kind)
        ws.column_dimensions[get_column_letter(first_col + j)].width = display_width(column)
    data = rows if rows else [{}]
    for i, record in enumerate(data):
        r = FIRST_ROW + i
        for j, column in enumerate(spec.columns):
            cell = ws.cell(row=r, column=first_col + j)
            if column.kind == "formula" and column.excel is not None:
                cell.value = W.formula_for(spec, column)
                cell.fill = FILL_FORMULA
                cell.font = FORMULA_FONT
            elif column.kind == "formula":
                cell.value = record.get("__formula__", {}).get(column.name)
                cell.fill = FILL_FORMULA
                cell.font = FORMULA_FONT
            else:
                cell.value = cell_value(record.get(column.name))
                cell.font = BODY
                if column.kind == "mirror":
                    cell.fill = FILL_MIRROR
            if column.fmt:
                cell.number_format = column.fmt
    last_row = FIRST_ROW + len(data) - 1
    last_col = first_col + len(spec.columns) - 1
    ref = f"{get_column_letter(first_col)}{HEADER_ROW}:{get_column_letter(last_col)}{last_row}"
    table = Table(displayName=spec.name, ref=ref)
    # Declare calculated columns so a row added in the spreadsheet fills in its formulas.
    table.tableColumns = []
    for j, column in enumerate(spec.columns):
        tc = TableColumn(id=j + 1, name=column.name)
        if column.kind == "formula" and column.excel is not None:
            tc.calculatedColumnFormula = TableFormula(attr_text=W.formula_for(spec, column)[1:])
        table.tableColumns.append(tc)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False, showColumnStripes=False)
    ws.add_table(table)
    return last_row, last_col


# Organization-specific ID lists warn instead of blocking, so an adopter can use their own role and profile IDs.
WARN_ONLY = {"role_id", "profile_id"}


def add_validation(ws, column_letter: str, lookup_type: str, start: int) -> None:
    dv = DataValidation(type="list", formula1=f"=lst_{lookup_type}", allow_blank=True, showErrorMessage=True,
                        errorStyle="warning" if lookup_type in WARN_ONLY else "stop",
                        errorTitle="Value not in R1 list", error=f"Choose a value from the R1 {lookup_type} list (Lookups tab).")
    dv.add(f"{column_letter}{start}:{column_letter}{start + VALIDATION_ROWS}")
    ws.add_data_validation(dv)


def build(output: Path, inputs: dict | None = None, settings: dict | None = None) -> dict:
    if openpyxl_xml.LXML:
        raise RuntimeError("openpyxl was imported with lxml serialization enabled; set OPENPYXL_LXML=False "
                           "before importing openpyxl so the build is byte-reproducible")
    ctx = RM.initial_context(inputs, settings)
    specs = W.tables()
    st = ctx["settings"]
    wb = Workbook()
    wb.remove(wb.active)
    sheets = {name: wb.create_sheet(name) for name in W.SHEETS}
    for name, ws in sheets.items():
        ws.sheet_view.zoomScale = 100
        ws["A1"] = SHEET_TITLES[name]
        ws["A1"].font = TITLE
        if name in SHEET_NOTES:
            ws["A2"] = SHEET_NOTES[name]
            ws["A2"].font = NOTE if name not in ("Escalation Rules", "Readiness Scorecard") else Font(name=FONT, size=11, bold=True, color="C00000")

    by_sheet: dict[str, list[W.TableSpec]] = {}
    for spec in specs.values():
        by_sheet.setdefault(spec.sheet, []).append(spec)

    defined: list[DefinedName] = []
    layout: dict[str, dict] = {}
    for sheet_name, sheet_specs in by_sheet.items():
        ws = sheets[sheet_name]
        start = 3 if sheet_name == "README" else 1
        positions = place_tables(sheet_specs, start)
        for spec in sheet_specs:
            rows = ctx.get(spec.name, [])
            last_row, last_col = write_table(ws, spec, rows, positions[spec.name], wb, defined)
            layout[spec.name] = {"sheet": sheet_name, "first_col": positions[spec.name], "last_row": last_row, "last_col": last_col}

    # Summary tables whose formulas are per metric row rather than per column.
    for spec_name, rows_def in (("tblKpiSummary", W.KPI_ROWS), ("tblCapacitySummary", W.CAPACITY_SUMMARY_ROWS)):
        ws = sheets[specs[spec_name].sheet]
        value_col = layout[spec_name]["first_col"] + [c.name for c in specs[spec_name].columns].index("value")
        for i, item in enumerate(rows_def):
            key = item[0]
            excel = W.kpi_excel(key, item[2]) if spec_name == "tblKpiSummary" else item[1]
            cell = ws.cell(row=FIRST_ROW + i, column=value_col, value="=" + excel)
            cell.fill = FILL_FORMULA
            cell.font = FORMULA_FONT

    # README text block and settings named cells.
    ws = sheets["README"]
    ws.column_dimensions["A"].width = 100
    ws.column_dimensions["B"].width = 2
    r = 3
    for kind, text in readme_lines(st):
        cell = ws.cell(row=r, column=1, value=text)
        cell.font = SUBTITLE if kind == "h" else BODY
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "p":
            ws.row_dimensions[r].height = 14 * max(1, -(-len(text) // 105))
        r += 1
    ws.cell(row=r + 1, column=1, value="Legend: dark blue header = you edit; grey header and grey cells = formula; green header = read-only mirror of R1.").font = NOTE
    settings_layout = layout["tblWorkbookSettings"]
    value_col = settings_layout["first_col"] + 1
    for i, row in enumerate(ctx["tblWorkbookSettings"]):
        cell = ws.cell(row=FIRST_ROW + i, column=value_col)
        if row["editable"] == "yes":
            cell.protection = Protection(locked=False)
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
        name = {"operating_mode": "operating_mode"}.get(row["setting_key"], row["setting_key"])
        if row["setting_key"] in ("calculation_as_of_at", "event_reconciliation_from_at", "selected_org_stage", "operating_mode"):
            defined.append(DefinedName(name, attr_text=f"'README'!${get_column_letter(value_col)}${FIRST_ROW + i}"))
    for i, row in enumerate(ctx["tblWorkbookSettings"]):
        if row["setting_key"] == "selected_org_stage":
            add_validation(ws, get_column_letter(value_col), "org_stage", FIRST_ROW + i)
        if row["setting_key"] == "operating_mode":
            add_validation(ws, get_column_letter(value_col), "operating_mode", FIRST_ROW + i)

    # Named lists from Lookups, and rule parameters.
    lk_layout = layout["tblLookupEntries"]
    key_col = get_column_letter(lk_layout["first_col"] + 2)
    entries = ctx["tblLookupEntries"]
    types: dict[str, list[int]] = {}
    for i, e in enumerate(entries):
        types.setdefault(e["lookup_type"], []).append(FIRST_ROW + i)
    for t_name, rows_idx in types.items():
        defined.append(DefinedName(f"lst_{t_name}", attr_text=f"'Lookups'!${key_col}${rows_idx[0]}:${key_col}${rows_idx[-1]}"))
    rp = layout["tblRuleParameters"]
    for i, p in enumerate(ctx["tblRuleParameters"]):
        if p["parameter"] != "readiness_calculation_version":
            defined.append(DefinedName(p["parameter"], attr_text=f"'Lookups'!${get_column_letter(rp['first_col'] + 1)}${FIRST_ROW + i}"))
    for d in defined:
        wb.defined_names[d.name] = d

    # Data validation on record tabs.
    for spec in specs.values():
        if spec.sheet in ("README", "Lookups", "Escalation Rules", "KPI Summary"):
            continue
        ws = sheets[spec.sheet]
        lay = layout[spec.name]
        for j, column in enumerate(spec.columns):
            if column.kind != "input":
                continue
            lookup = column.validation or FIELD_VALIDATION.get(column.name)
            if column.name == "status":
                lookup = STATUS_VALIDATION.get(spec.name)
            if column.name.endswith("_role_id"):
                lookup = "role_id"
            letter = get_column_letter(lay["first_col"] + j)
            if lookup:
                add_validation(ws, letter, lookup, FIRST_ROW)
            if spec.name == "tblRisks" and column.name in ("likelihood", "impact"):
                dv = DataValidation(type="whole", operator="between", formula1=f"=risk_{column.name}_min", formula2=f"=risk_{column.name}_max",
                                    allow_blank=True, showErrorMessage=True, errorTitle="Outside the R1 scale",
                                    error="Use a whole number on the R1 risk scale (Lookups, tblRuleParameters).")
                dv.add(f"{letter}{FIRST_ROW}:{letter}{FIRST_ROW + VALIDATION_ROWS}")
                ws.add_data_validation(dv)

    # Conditional formatting for scanning.
    for spec in specs.values():
        ws = sheets[spec.sheet]
        lay = layout[spec.name]
        for j, column in enumerate(spec.columns):
            if column.kind != "formula":
                continue
            letter = get_column_letter(lay["first_col"] + j)
            rng = f"{letter}{FIRST_ROW}:{letter}{FIRST_ROW + VALIDATION_ROWS}"
            top = f"${letter}{FIRST_ROW}"
            if column.name == "event_evidence":
                ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({top},7)="missing"'], fill=FILL_WARN))
            elif column.name in ("reference_check", "gate_check", "status_check", "severity_check", "outcome_check", "category_count_check",
                                 "escalation_due_check", "predecessor_reference_valid", "subject_check", "actor_check", "event_type_check"):
                ok = {"category_count_check": "complete", "escalation_due_check": "consistent", "subject_check": "found",
                      "event_type_check": "registered", "predecessor_reference_valid": "valid"}.get(column.name, "valid")
                extra = ',{0}<>"no_predecessor",{0}<>"event_only_record"'.format(top)
                ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND({top}<>"",{top}<>"{ok}"{extra})'], fill=FILL_BAD))
            elif column.name == "due_state":
                ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{top}="overdue"'], fill=FILL_BAD))
                ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{top}="blocked"'], fill=FILL_WARN))
            elif column.name == "escalation_due_state":
                ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{top}="escalation_due"'], fill=FILL_WARN))
            elif column.name == "evidence_warning":
                ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({top},28)="Required evidence incomplete"'], fill=FILL_WARN))

    # Chart: Projects by Current Phase.
    ps = layout["tblPhaseSummary"]
    ws = sheets["KPI Summary"]
    chart = BarChart()
    chart.type = "bar"
    chart.title = "Projects by Current Phase"
    chart.y_axis.title = "Projects"
    chart.y_axis.majorUnit = 1  # whole projects, never fractional ticks
    chart.y_axis.scaling.min = 0
    chart.y_axis.number_format = "0"
    chart.legend = None
    data = Reference(ws, min_col=ps["first_col"] + 2, min_row=HEADER_ROW, max_row=ps["last_row"])
    cats = Reference(ws, min_col=ps["first_col"] + 1, min_row=FIRST_ROW, max_row=ps["last_row"])
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height, chart.width = 9, 16
    ws.add_chart(chart, f"A{FIRST_ROW + len(W.KPI_ROWS) + 3}")

    # Freeze panes, filters (tables carry their own), and protection.
    for name, ws in sheets.items():
        if name != "README":
            ws.freeze_panes = f"A{FIRST_ROW}"
    for name in ("README", "Lookups", "Escalation Rules", "KPI Summary"):
        sheets[name].protection.sheet = True
        sheets[name].protection.autoFilter = False
        sheets[name].protection.sort = False
        sheets[name].protection.formatColumns = False

    when = datetime.strptime(st["generated_at"], "%Y-%m-%dT%H:%M:%SZ")
    props = wb.properties
    props.creator = "Jared Muldrow"
    props.lastModifiedBy = "Jared Muldrow"
    props.title = "Implementation Tracker Workbook"
    props.subject = "Macro-free reference implementation of the R1 implementation operating model"
    props.keywords = "implementation; tracker; synthetic data; illustrative example"
    props.created = when
    props.modified = when
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    output.parent.mkdir(parents=True, exist_ok=True)
    buffer = io.BytesIO()
    wb.save(buffer)
    output.write_bytes(normalize_zip(buffer.getvalue(), st["generated_at"]))
    return layout


def normalize_zip(data: bytes, generated_at: str) -> bytes:
    """Rewrite the package with fixed entry times, fixed document timestamps, and no
    compression, so a rebuild from the same inputs gives byte-identical output on any
    machine (compressed bytes can differ between zlib versions)."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_STORED) as dst:
        for info in src.infolist():
            content = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                text = content.decode("utf-8")
                text = re.sub(r"(<dcterms:(created|modified)[^>]*>)[^<]*(</dcterms:\2>)", rf"\g<1>{generated_at}\g<3>", text)
                content = text.encode("utf-8")
            entry = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_STORED
            entry.external_attr = 0o600 << 16
            dst.writestr(entry, content)
    return out.getvalue()


# ---------------------------------------------------------------------------
# Recalculation and exports


def recalculate(xlsx: Path, out_dir: Path) -> Path:
    """Recalculate a copy with headless LibreOffice. Raises FileNotFoundError when
    LibreOffice is not installed. The source file is never modified."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise FileNotFoundError("LibreOffice (soffice) is not installed")
    out_dir.mkdir(parents=True, exist_ok=True)
    profile = Path(tempfile.mkdtemp(prefix="lo-profile-"))
    try:
        subprocess.run([soffice, f"-env:UserInstallation=file://{profile}", "--headless", "--calc", "--convert-to", "xlsx",
                        "--outdir", str(out_dir), str(xlsx)], capture_output=True, check=True, timeout=300)
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    result = out_dir / xlsx.name
    if not result.is_file():
        raise RuntimeError("LibreOffice produced no recalculated file")
    return result


def read_tables(xlsx: Path, data_only: bool = True) -> dict[str, list[dict]]:
    wb = load_workbook(xlsx, data_only=data_only)
    out: dict[str, list[dict]] = {}
    for ws in wb.worksheets:
        for name, table in ws.tables.items():
            ref = table.ref if hasattr(table, "ref") else table
            rows = list(ws[ref])
            header = [c.value for c in rows[0]]
            out[name] = [{h: c.value for h, c in zip(header, row)} for row in rows[1:]]
    return out


def export_value(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return f"{v:.10f}".rstrip("0").rstrip(".")
    return str(v)


def write_exports(recalculated: Path, root: Path) -> list[Path]:
    export_map = W.yaml.safe_load((W.ROOT / "config" / "export-map.yaml").read_text(encoding="utf-8"))
    tables = read_tables(recalculated)
    written = []
    (root / "exports" / "csv").mkdir(parents=True, exist_ok=True)
    (root / "exports" / "jsonl").mkdir(parents=True, exist_ok=True)
    for item in export_map["exports"]:
        rows = [r for r in tables[item["table"]] if r.get(item["columns"][0]) not in (None, "")]
        path = root / item["file"]
        if path.suffix == ".csv":
            with path.open("w", encoding="utf-8", newline="") as h:
                w = csv.writer(h, lineterminator="\n")
                w.writerow(item["columns"])
                for r in rows:
                    w.writerow([export_value(r.get(c)) for c in item["columns"]])
        else:
            with path.open("w", encoding="utf-8") as h:
                for r in rows:
                    record = {c: r.get(c) for c in item["columns"]}
                    record["payload"] = json.loads(record["payload"])
                    h.write(json.dumps(record) + "\n")
        written.append(path)
    return written


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build the tracker workbook reproducibly.")
    parser.add_argument("--output", default=str(W.WORKBOOK_PATH))
    parser.add_argument("--data-dir", default=None, help="Alternative R1-format synthetic data folder (tests)")
    parser.add_argument("--capacity", default=None, help="Alternative capacity-input definitions CSV (tests)")
    parser.add_argument("--empty", action="store_true", help="Build with no records (tests)")
    parser.add_argument("--exports", action="store_true", help="Also recalculate with LibreOffice and write exports/")
    parser.add_argument("--docs", action="store_true", help="Regenerate FORMULAS.md and data-dictionary inventories and the expected values")
    args = parser.parse_args(argv)
    if args.docs:
        write_docs()
        write_expected_values()
        print("regenerated FORMULAS.md, docs/data-dictionary.md, and verification/expected-values.yaml")
        return 0
    inputs = None
    if args.empty:
        inputs = {k: [] for k in W.load_inputs()}
    elif args.data_dir or args.capacity:
        inputs = W.load_inputs(Path(args.data_dir) if args.data_dir else None, Path(args.capacity) if args.capacity else None)
    out = Path(args.output)
    build(out, inputs)
    print(f"built {out}")
    if args.exports:
        with tempfile.TemporaryDirectory() as tmp:
            recalculated = recalculate(out, Path(tmp))
            for p in write_exports(recalculated, W.ROOT):
                print(f"exported {p.relative_to(W.ROOT)}")
    return 0



# ---------------------------------------------------------------------------
# Generated documentation (FORMULAS.md tables and docs/data-dictionary.md)

R1_SOURCES = {
    "risk_score": "config/risk-rules.yaml (formula likelihood_times_impact; likelihood_values, impact_values)",
    "risk_band": "config/risk-rules.yaml (bands)", "risk_severity": "config/risk-rules.yaml (bands.severity)",
    "weight": "config/readiness-weights.yaml (categories.weight)", "achieved_rate": "config/readiness-weights.yaml (calculation)",
    "weighted_result": "config/readiness-weights.yaml (calculation)", "category_score": "config/readiness-weights.yaml (calculation)",
    "overall_score": "config/readiness-weights.yaml (calculation)", "incomplete_required_evidence_flag": "config/readiness-weights.yaml (calculation)",
    "escalation_rule_id": "config/sla-rules.yaml (rules.trigger, rules.action)", "escalation_due_check": "config/sla-rules.yaml (rules.target_duration)",
    "gate_check": "lifecycle/gates.yaml (gates.phase_id)", "outcome_check": "standard/lifecycle-terms.yaml (gate_outcomes)",
    "assessor_check": "schemas/gate-assessment.schema.json (assessor_person_id)", "acceptance_check": "schemas/handoff.schema.json (accepted status rule)",
    "severity_check": "standard/severity-scale.yaml", "event_type_check": "standard/event-catalog.yaml",
    "source_repo_check": "standard/event-catalog.yaml (producer_repo)", "actor_check": "standard/schemas/event.schema.json (actor rules)",
    "request_transition_check": "config/request-state-machine.yaml (transitions)",
}
STATUS_SOURCES = {"tblRequests": "config/request-state-machine.yaml (statuses)", "tblHandoffs": "schemas/handoff.schema.json (status)"}
KPI_INTERPRETATION = {
    "project_count": "How many projects the workbook tracks.",
    "active_project_count": "Projects in delivery now.",
    "open_task_count": "Work not yet finished, including blocked and not started tasks.",
    "blocked_task_count": "Tasks waiting on something outside the task; look for the matching request or issue.",
    "overdue_task_count": "Open tasks past their planned due date as of calculation_as_of_at; re-plan or escalate.",
    "open_risk_count": "Risks still being worked; accepted and closed risks are excluded.",
    "open_issue_count": "Problems affecting delivery now.",
    "request_count": "All requests in the workbook.",
    "escalation_due_request_count": "Requests whose R1 escalation time has passed; the escalation role should already be involved.",
    "readiness_scorecard_count": "Scorecards recorded, including early pre-assessments.",
    "incomplete_readiness_evidence_count": "Scorecards with missing required evidence; a person resolves these before any launch decision.",
    "event_log_row_count": "Rows in the Event Log.",
    "event_log_problem_count": "Event rows that fail a check; must be 0 before an Event Log export is trusted.",
    "missing_event_evidence_count": "Material changes inside the reconciliation window that still need an Event Log row.",
    "reference_problem_count": "Records whose project, phase, or milestone reference does not resolve; must be 0.",
}


def _refs(formula: str) -> list[str]:
    seen = []
    for m in re.findall(r"\[\[#This Row\],\[([a-z0-9_]+)\]\]|tbl[A-Za-z]+\[([a-z0-9_]+)\]", formula):
        name = m[0] or m[1]
        if name and name not in seen:
            seen.append(name)
    for named in ("calculation_as_of_at", "event_reconciliation_from_at", "risk_likelihood_min", "risk_likelihood_max",
                  "risk_impact_min", "risk_impact_max", "lst_phase_key", "lst_request_status"):
        if named in formula and named not in seen:
            seen.append(named)
    return seen


def _result_summary(values: list) -> str:
    values = [("" if v is None else v) for v in values]
    if not values:
        return "no synthetic rows"
    if all(isinstance(v, bool) for v in values):
        return f"TRUE {sum(values)}, FALSE {len(values) - sum(values)}"
    if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in values) and len(set(values)) > 6:
        return f"{len(values)} values, total {round(sum(values), 4):g}"
    counts: dict = {}
    for v in values:
        key = "(blank)" if v == "" else (f"{v:g}" if isinstance(v, float) else str(v))
        counts[key] = counts.get(key, 0) + 1
    return ", ".join(f"{k} {n}" for k, n in counts.items())


def _portability(formula: str) -> str:
    used = sorted(W.formula_names(formula))
    note = "Functions: " + ", ".join(used) + " (all in the portable set)." if used else "No functions."
    if "[#This Row]" in formula or "tbl" in formula:
        note += " Uses Excel Table structured references; recalculation verified in LibreOffice Calc 24.2, not yet in Excel itself."
    if "DATEVALUE" in formula:
        note += " DATEVALUE and TIMEVALUE read ISO text; locales that do not accept ISO dates are a documented limitation."
    return note


def formulas_markdown(ctx: dict) -> str:
    specs = W.tables()
    out = ["<!-- Generated by tools/build_workbook.py --docs from tools/workbook_spec.py. Do not edit by hand. -->", ""]
    for spec in specs.values():
        cols = [c for c in spec.columns if c.kind == "formula" and c.excel is not None]
        if not cols:
            continue
        out.append(f"### {spec.sheet} / `{spec.name}`")
        out.append("")
        for c in cols:
            formula = W.formula_for(spec, c)
            source = R1_SOURCES.get(c.name) or (STATUS_SOURCES.get(spec.name) if c.name == "status_check" else None) or "None: R3 view over R1 records"
            out += [
                f"#### `{spec.name}[{c.name}]`", "",
                "| Item | Value |", "|---|---|",
                f"| Workbook tab | {spec.sheet} |", f"| Table | `{spec.name}` |", f"| Field | `{c.name}` ({'R1 field' if c.owner == 'R1' else 'R3 helper, not exported' if not c.export else 'R3 field'}) |",
                f"| Authority | `{c.authority}` |", f"| Purpose | {c.purpose} |", f"| R1 source rule or config | {source} |",
                f"| Input fields | {', '.join(f'`{r}`' for r in _refs(formula)) or 'none'} |", f"| Edge cases | {c.edge_cases} |",
                f"| Expected synthetic result | {_result_summary([r.get(c.name) for r in ctx[spec.name]])} |",
                f"| Portability | {_portability(formula)} |", "",
                "```text", formula, "```", "",
            ]
    out += ["### KPI Summary / `tblKpiSummary`", "",
            "| KPI | Definition | Source table | Grain | Interpretation | Authority | Expected synthetic value |", "|---|---|---|---|---|---|---|"]
    for (key, label, excel, _, grain, src, auth), row in zip(W.KPI_ROWS, ctx["tblKpiSummary"]):
        out.append(f"| `{key}` | {label} | `{src}` | {grain} | {KPI_INTERPRETATION[key]} | `{auth}` | {row['value']} |")
    out += ["", "KPI formulas:", ""]
    for key, label, excel, *_ in W.KPI_ROWS:
        out += [f"- `{key}`", "", "  ```text", "  =" + W.kpi_excel(key, excel), "  ```", ""]
    out += ["### Capacity Inputs / `tblCapacitySummary`", "",
            "| Check | Meaning | Expected synthetic value |", "|---|---|---|"]
    for (key, excel, _, meaning), row in zip(W.CAPACITY_SUMMARY_ROWS, ctx["tblCapacitySummary"]):
        out.append(f"| `{key}` | {meaning} | {row['value']} |")
    out += ["", "Capacity check formulas:", ""]
    for key, excel, *_ in W.CAPACITY_SUMMARY_ROWS:
        out += [f"- `{key}`", "", "  ```text", "  =" + excel, "  ```", ""]
    return "\n".join(out).rstrip() + "\n"


def data_dictionary_markdown() -> str:
    specs = W.tables()
    lines = ["<!-- Generated by tools/build_workbook.py --docs from tools/workbook_spec.py. Do not edit by hand. -->", ""]
    export_map = W.yaml.safe_load((W.ROOT / "config" / "export-map.yaml").read_text(encoding="utf-8"))
    exported = {(e["table"], c): e["file"] for e in export_map["exports"] for c in e["columns"]}
    for spec in specs.values():
        lines += [f"## {spec.sheet} / `{spec.name}`", "", f"{spec.title}. Entity: `{spec.entity}` ({spec.owner}).", "",
                  "| Display label | Authoritative field name | Entity | Type | Edit mode | Required | Ownership | Validation source | Export |",
                  "|---|---|---|---|---|---|---|---|---|"]
        schema = None
        if spec.owner == "R1" and spec.entity not in ("readiness_weight", "risk_band", "rule_parameter", "rule", "event"):
            schema = W.schema_for(spec.entity)
        elif spec.entity == "event":
            schema = W.load_json("standard/schemas/event.schema.json")
        elif spec.entity in ("capacity-input",):
            schema = json.loads((W.ROOT / "schemas/r3/capacity-input.schema.json").read_text(encoding="utf-8"))
        for c in spec.columns:
            ptype = ""
            if schema:
                p = schema["properties"].get(c.name) or schema["properties"].get("categories", {}).get("items", {}).get("properties", {}).get(c.name) or {}
                ptype = p.get("type", "")
                ptype = "/".join(ptype) if isinstance(ptype, list) else ptype
                if p.get("format") in ("date", "date-time"):
                    ptype = f"{p['format']} (ISO text)"
            ptype = ptype or ("text" if c.kind != "formula" else "calculated")
            mode = {"input": "editable", "formula": "formula", "mirror": "read-only mirror"}[c.kind]
            req = "yes" if (c.required or (schema and c.name in schema.get("required", []))) else "no"
            owner = "R1 field" if c.owner == "R1" else "R3"
            lookup = c.validation or FIELD_VALIDATION.get(c.name)
            if c.name == "status":
                lookup = STATUS_VALIDATION.get(spec.name)
            if c.name.endswith("_role_id") and c.kind == "input":
                lookup = "role_id"
            vsrc = f"`lst_{lookup}`" if lookup and c.kind == "input" else ("pinned R1 configuration" if c.kind == "mirror" else "none")
            exp = exported.get((spec.name, c.name), "not exported")
            label = c.name.replace("_", " ").capitalize()
            lines.append(f"| {label} | `{c.name}` | {spec.entity} | {ptype} | {mode} | {req} | {owner} | {vsrc} | {exp} |")
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def write_expected_values() -> None:
    values = json.loads(json.dumps(RM.expected_values()))
    header = ("# Expected workbook values for the bundled synthetic data. Generated by tools/build_workbook.py --docs\n"
              "# from the pinned R1 inputs; tools/verify_workbook.py checks the recalculated workbook against it.\n"
              "# All values are synthetic data, not measured results.\n")
    (W.ROOT / "verification" / "expected-values.yaml").write_text(header + W.yaml.safe_dump(values, sort_keys=False, width=160), encoding="utf-8")


def write_docs() -> None:
    ctx = RM.compute()
    formulas = (W.ROOT / "FORMULAS.md").read_text(encoding="utf-8")
    marker = "<!-- GENERATED FORMULA INVENTORY BELOW -->"
    head = formulas.split(marker)[0] if marker in formulas else formulas
    (W.ROOT / "FORMULAS.md").write_text(head.rstrip() + "\n\n" + marker + "\n\n" + formulas_markdown(ctx), encoding="utf-8")
    dd = (W.ROOT / "docs" / "data-dictionary.md").read_text(encoding="utf-8")
    marker2 = "<!-- GENERATED DATA DICTIONARY BELOW -->"
    head2 = dd.split(marker2)[0] if marker2 in dd else dd
    (W.ROOT / "docs" / "data-dictionary.md").write_text(head2.rstrip() + "\n\n" + marker2 + "\n\n" + data_dictionary_markdown(), encoding="utf-8")


if __name__ == "__main__":
    sys.exit(main())
