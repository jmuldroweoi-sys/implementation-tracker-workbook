"""Inspect the generated workbook and verify it against the pinned R1 version.

Structural checks (no spreadsheet engine needed):
  sheet names and order, no hidden sheets, required tables and headers, formulas where
  expected and only there, no macros, no external links, no volatile or unapproved
  functions, validation lists equal to pinned R1 vocabularies, metadata equal to the pin,
  synthetic inputs equal to the pinned R1 records, a valid workbook-metadata record, and a
  byte-identical rebuild.

Recalculation checks (headless LibreOffice, when installed):
  the workbook is recalculated by a real spreadsheet engine; no cell holds an error value;
  every formula cell equals the reference evaluator's value; the expected values file is
  current; and exports regenerated from the recalculated tables equal the committed exports.

If LibreOffice is not installed, recalculation checks are reported as NOT RUN, never as
passed. Use --require-recalc (CI) to make NOT RUN a failure.

Usage:
    python tools/verify_workbook.py [--workbook PATH] [--require-recalc] [--skip-rebuild]
"""

from __future__ import annotations

import argparse
import filecmp
import json
import re
import sys
import tempfile
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import yaml  # noqa: E402
from jsonschema import Draft202012Validator  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import build_workbook as B  # noqa: E402
import reference_model as RM  # noqa: E402
import workbook_spec as W  # noqa: E402

ERROR_VALUES = ("#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!", "#SPILL!", "#CALC!", "Err:")


class Report:
    def __init__(self) -> None:
        self.rows: list[tuple[str, str, str]] = []

    def add(self, check: str, problems: list[str], note: str = "", not_run: bool = False) -> None:
        status = "NOT RUN" if not_run else ("FAIL" if problems else "PASS")
        self.rows.append((check, status, "; ".join(problems[:6]) if problems else note))

    @property
    def failed(self) -> bool:
        return any(s == "FAIL" for _, s, _ in self.rows)


def same(a, b) -> bool:
    a = "" if a is None else a
    b = "" if b is None else b
    if isinstance(a, bool) or isinstance(b, bool):
        return a is b or (a == b and type(a) is type(b))
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) < 1e-9
    return a == b


def formula_cells(xlsx: Path) -> list[tuple[str, str, str]]:
    wb = load_workbook(xlsx)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    out.append((ws.title, c.coordinate, c.value))
    return out


def structural_checks(xlsx: Path, report: Report, inputs: dict | None = None) -> None:
    specs = W.tables()
    wb = load_workbook(xlsx)
    names = wb.sheetnames
    report.add("W01 exact sheet names and order", [] if names == list(W.SHEETS) else [f"sheets are {names}"], f"{len(names)} sheets in order")
    hidden = [ws.title for ws in wb.worksheets if ws.sheet_state != "visible"]
    report.add("W02 no hidden sheets", [f"hidden: {h}" for h in hidden], "12 visible, 0 hidden")

    tables = {}
    for ws in wb.worksheets:
        for table in ws.tables.values():
            tables[table.displayName] = (ws, table.ref)
    problems = []
    for name, spec in specs.items():
        if name not in tables:
            problems.append(f"missing table {name}")
            continue
        ws, ref = tables[name]
        if ws.title != spec.sheet:
            problems.append(f"{name} is on {ws.title}, expected {spec.sheet}")
        header = [c.value for c in list(ws[ref])[0]]
        if header != [c.name for c in spec.columns]:
            problems.append(f"{name}: headers differ from the specification")
    problems += [f"unexpected table {n}" for n in tables if n not in specs]
    report.add("W03 required tables and headers", problems, f"{len(specs)} tables")

    problems = []
    count = 0
    for name, spec in specs.items():
        if name not in tables:
            continue
        ws, ref = tables[name]
        rows = list(ws[ref])[1:]
        for j, column in enumerate(spec.columns):
            for row in rows:
                v = row[j].value
                is_formula = isinstance(v, str) and v.startswith("=")
                expect_formula = column.kind == "formula"
                if expect_formula and not is_formula:
                    problems.append(f"{name}.{column.name} {row[j].coordinate}: formula missing")
                if not expect_formula and is_formula:
                    problems.append(f"{name}.{column.name} {row[j].coordinate}: unexpected formula in an input cell")
                if is_formula:
                    count += 1
                    expected = W.formula_for(spec, column) if column.excel else None
                    if expected and v != expected:
                        problems.append(f"{name}.{column.name} {row[j].coordinate}: formula differs from the specification")
    report.add("W04 formulas present where expected", problems, f"{count} formula cells in record tables")

    with zipfile.ZipFile(xlsx) as z:
        members = z.namelist()
        content_types = z.read("[Content_Types].xml").decode("utf-8")
        rels = "".join(z.read(m).decode("utf-8", "replace") for m in members if m.endswith(".rels"))
    macro = [m for m in members if "vbaProject" in m or m.lower().endswith(".bin")]
    if "macroEnabled" in content_types:
        macro.append("macro-enabled content type")
    report.add("W05 no macros", macro, "no VBA project, not macro-enabled")
    ext = [m for m in members if "externalLink" in m]
    if 'TargetMode="External"' in rels:
        ext.append("external relationship target")
    all_formulas = formula_cells(xlsx)
    ext += [f"{s}!{c}" for s, c, f in all_formulas if re.search(r"\[[^\]]*\.xls|https?://", f)]
    report.add("W06 no external links", ext, "no external link parts, targets, or file references")
    bad = []
    for s, c, f in all_formulas:
        used = W.formula_names(f)
        for fn in W.FORBIDDEN_FUNCTIONS:
            if fn in used:
                bad.append(f"{s}!{c} uses {fn}")
    report.add("W07 no volatile or unapproved functions", bad, f"{len(all_formulas)} formulas scanned; TODAY, NOW, OFFSET, INDIRECT, LET, LAMBDA, XLOOKUP absent")

    problems = []
    entries = W.lookup_entries()
    for src in W.LOOKUP_SOURCES:
        dn = wb.defined_names.get(f"lst_{src.lookup_type}")
        if dn is None:
            problems.append(f"named list lst_{src.lookup_type} missing")
            continue
        sheet, rng = list(dn.destinations)[0]
        values = [c.value for row in wb[sheet][rng.replace("$", "")] for c in row]
        if values != W.lookup_values(entries, src.lookup_type):
            problems.append(f"lst_{src.lookup_type} differs from the pinned source {src.source_path}")
    lk_ws, lk_ref = tables["tblLookupEntries"]
    rows = [[c.value for c in r] for r in list(lk_ws[lk_ref])[1:]]
    want = [[e[k] for k in ("lookup_entry_id", "lookup_type", "key", "display_value", "source_repo", "source_path", "source_version", "active_flag")] for e in entries]
    if len(rows) != len(want) or not all(same(a, b) for g, w in zip(rows, want) for a, b in zip(g, w)):
        problems.append("tblLookupEntries differs from the pinned R1 sources")
    for mirror, rows_expected in (("tblReadinessWeights", W.readiness_weights()), ("tblRiskBands", W.risk_bands()),
                                  ("tblRuleParameters", W.rule_parameters()), ("tblEscalationRules", W.escalation_rules())):
        ws, ref = tables[mirror]
        got = [[c.value for c in r] for r in list(ws[ref])[1:]]
        cols = [c.name for c in specs[mirror].columns]
        want = [[e[c] for c in cols] for e in rows_expected]
        if len(got) != len(want) or not all(same(a, b) for g, w in zip(got, want) for a, b in zip(g, w)):
            problems.append(f"{mirror} differs from the pinned R1 configuration")
    report.add("W08 validation lists and mirrors match R1", problems, f"{len(W.LOOKUP_SOURCES)} named lists and 4 mirror tables equal the pin")

    problems = []
    ref = W.pin_reference()
    st_ws, st_ref = tables["tblWorkbookSettings"]
    settings = {r[0].value: r[1].value for r in list(st_ws[st_ref])[1:]}
    if settings.get("r1_source_commit") != ref["r1_commit"]:
        problems.append("r1_source_commit differs from standard/standard-reference.yaml")
    if settings.get("shared_standard_version") != ref["shared_standard_version"]:
        problems.append("shared_standard_version differs from the pin")
    metadata = {"workbook_id": settings.get("workbook_id"), "workbook_version": settings.get("workbook_version"),
                "r1_commit": settings.get("r1_source_commit"), "shared_standard_version": settings.get("shared_standard_version"),
                "calculation_as_of_at": settings.get("calculation_as_of_at"), "event_reconciliation_from_at": settings.get("event_reconciliation_from_at"),
                "org_stage": settings.get("selected_org_stage"), "operating_mode": settings.get("operating_mode"), "generated_at": settings.get("generated_at")}
    schema = json.loads((W.ROOT / "schemas/r3/workbook-metadata.schema.json").read_text(encoding="utf-8"))
    problems += [f"metadata: {e.message}" for e in Draft202012Validator(schema).iter_errors(metadata)]
    if wb.properties.creator != "Jared Muldrow":
        problems.append("document creator is not Jared Muldrow")
    report.add("W09 metadata matches the pin and validates", problems, f"WBK metadata valid; R1 {ref['r1_commit'][:7]}; standard {ref['shared_standard_version']}")

    problems = []
    inputs = inputs if inputs is not None else W.load_inputs()
    for name, rows_in in inputs.items():
        spec = specs[name]
        ws, tref = tables[name]
        sheet_rows = list(ws[tref])[1:]
        if rows_in and len(sheet_rows) != len(rows_in):
            problems.append(f"{name}: {len(sheet_rows)} rows, source has {len(rows_in)}")
            continue
        for record, row in zip(rows_in, sheet_rows):
            for j, column in enumerate(spec.columns):
                if column.kind != "input":
                    continue
                want = B.cell_value(record.get(column.name))
                if not same(row[j].value, want):
                    problems.append(f"{name} {record.get(spec.id_column)}.{column.name}: workbook value differs from the source")
    report.add("W10 synthetic inputs equal the pinned sources", problems, "every input cell equals the pinned R1 record or R3 synthetic row")


def workbook_inputs(xlsx: Path) -> dict[str, list[dict]]:
    """Input rows of every record table, read from the workbook itself (formula cells are
    ignored). Values are converted to R1 field conventions: arrays split on semicolons,
    blank cells omitted."""
    tables = B.read_tables(xlsx, data_only=False)
    out: dict[str, list[dict]] = {}
    for name in W.load_inputs().keys():
        spec = W.tables()[name]
        schema = None
        if spec.owner == "R1" and spec.entity != "event":
            schema = W.schema_for(spec.entity)
        rows = []
        for raw in tables.get(name, []):
            if raw.get(spec.id_column) in (None, ""):
                continue
            rec = {}
            for column in spec.columns:
                if column.kind != "input":
                    continue
                v = raw.get(column.name)
                if v is None or v == "":
                    continue
                prop = {}
                if schema:
                    prop = schema["properties"].get(column.name) or schema["properties"].get("categories", {}).get("items", {}).get("properties", {}).get(column.name, {})
                if prop.get("type") == "array" and isinstance(v, str):
                    v = v.split(";")
                rec[column.name] = v
            rows.append(rec)
        out[name] = rows
    return out


def rebuild_check(xlsx: Path, report: Report) -> None:
    with tempfile.TemporaryDirectory() as tmp:
        again = Path(tmp) / "rebuild.xlsx"
        B.build(again)
        same_bytes = filecmp.cmp(again, xlsx, shallow=False)
    report.add("W11 reproducible build (byte-identical rebuild)", [] if same_bytes else ["rebuilt workbook differs from the committed file"],
               "rebuild from pinned inputs is byte-identical")


def recalculation_checks(xlsx: Path, report: Report, require: bool) -> None:
    labels = ("W12 recalculated: no error values", "W13 recalculated: every formula equals the reference evaluator",
              "W14 expected values current", "W15 exports regenerate from the recalculated workbook")
    with tempfile.TemporaryDirectory() as tmp:
        try:
            recalculated = B.recalculate(xlsx, Path(tmp) / "out")
        except FileNotFoundError:
            for label in labels:
                report.add(label, ["LibreOffice is required (--require-recalc)"] if require else [], "LibreOffice not installed", not_run=not require)
            return
        wb = load_workbook(recalculated, data_only=True)
        errors = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith(ERROR_VALUES):
                        errors.append(f"{ws.title}!{c.coordinate} {c.value}")
        report.add(labels[0], errors, "0 error values after headless LibreOffice recalculation")

        ctx = RM.compute()
        got = B.read_tables(recalculated)
        problems, compared = [], 0
        for name, spec in W.tables().items():
            for column in spec.columns:
                if column.kind != "formula":
                    continue
                for i, (g, e) in enumerate(zip(got[name], ctx[name])):
                    compared += 1
                    if not same(g.get(column.name), e.get(column.name)):
                        problems.append(f"{name}[{i + 1}].{column.name}: workbook {g.get(column.name)!r}, reference {e.get(column.name)!r}")
        report.add(labels[1], problems, f"{compared} formula values equal the reference evaluator")

        expected = yaml.safe_load((W.ROOT / "verification/expected-values.yaml").read_text(encoding="utf-8"))
        current = RM.expected_values(ctx)
        report.add(labels[2], [] if expected == json.loads(json.dumps(current)) else ["verification/expected-values.yaml is stale; regenerate it"],
                   "expected values equal the reference evaluator")

        out_root = Path(tmp) / "exports"
        problems = []
        for p in B.write_exports(recalculated, out_root):
            committed = W.ROOT / p.relative_to(out_root)
            if not committed.is_file() or committed.read_bytes() != p.read_bytes():
                problems.append(f"{committed.relative_to(W.ROOT)} differs from a fresh export")
        report.add(labels[3], problems, "13 exports byte-identical to a fresh export of the recalculated workbook")


def verify(xlsx: Path, require_recalc: bool = False, rebuild: bool = True) -> Report:
    report = Report()
    structural_checks(xlsx, report)
    if rebuild:
        rebuild_check(xlsx, report)
    recalculation_checks(xlsx, report, require_recalc)
    return report


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Verify the tracker workbook.")
    parser.add_argument("--workbook", default=str(W.WORKBOOK_PATH))
    parser.add_argument("--require-recalc", action="store_true")
    parser.add_argument("--skip-rebuild", action="store_true")
    args = parser.parse_args(argv)
    report = verify(Path(args.workbook), args.require_recalc, not args.skip_rebuild)
    for check, status, note in report.rows:
        print(f"{status} {check}: {note}")
    failed = sum(1 for _, s, _ in report.rows if s == "FAIL")
    not_run = sum(1 for _, s, _ in report.rows if s == "NOT RUN")
    print(f"RESULT: {'FAIL' if failed else 'PASS'} ({len(report.rows) - failed - not_run} passed, {failed} failed, {not_run} not run)")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
