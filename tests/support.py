"""Shared test helpers. Every test works on copies and temporary files; the committed
repository is never modified. All test data is synthetic data and an illustrative example."""

from __future__ import annotations

import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))

import build_workbook as B  # noqa: E402,F401
import reference_model as RM  # noqa: E402,F401
import workbook_spec as W  # noqa: E402

REQUIRE_RECALC = os.environ.get("R3_REQUIRE_RECALC") == "1"


def libreoffice_available() -> bool:
    return bool(shutil.which("soffice") or shutil.which("libreoffice"))


def needs_libreoffice(test):
    """Skip without LibreOffice locally; in CI (R3_REQUIRE_RECALC=1) a missing engine fails."""
    if libreoffice_available() or REQUIRE_RECALC:
        return test
    return unittest.skip("LibreOffice not installed; recalculation tests run in CI")(test)


def reset_root() -> None:
    W.ROOT = ROOT
    W.PIN = ROOT / "schemas" / "r1"
    W.WORKBOOK_PATH = ROOT / "workbook" / "implementation-tracker-workbook.xlsx"
    W.TABLES.clear()


class Sandbox:
    """A full copy of the repository (without .git) in a temporary folder."""

    def __init__(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name) / "repo"
        shutil.copytree(ROOT, self.root, ignore=shutil.ignore_patterns(".git", "__pycache__"))

    def path(self, rel: str) -> Path:
        return self.root / rel

    def replace(self, rel: str, old: str, new: str) -> None:
        text = self.path(rel).read_text(encoding="utf-8")
        assert old in text, old
        self.path(rel).write_text(text.replace(old, new, 1), encoding="utf-8")

    def edit_workbook(self, change) -> None:
        from openpyxl import load_workbook

        p = self.path("workbook/implementation-tracker-workbook.xlsx")
        wb = load_workbook(p)
        change(wb)
        wb.save(p)

    def run(self, blocklist=None) -> dict[str, tuple[bool, str]]:
        import validate

        try:
            res = validate.validate(self.root, blocklist)
        finally:
            reset_root()
        return {check.split(" ")[0]: (ok, note) for check, ok, note in res.rows}

    def close(self) -> None:
        self.tmp.cleanup()
        reset_root()


def table_cell(wb, table: str, row_id: str, column: str):
    """Return the worksheet cell of a table row (found by its first-column ID)."""
    for ws in wb.worksheets:
        for t in ws.tables.values():
            if t.displayName != table:
                continue
            rows = list(ws[t.ref])
            header = [c.value for c in rows[0]]
            j = header.index(column)
            for r in rows[1:]:
                if r[0].value == row_id:
                    return r[j]
    raise KeyError((table, row_id, column))
