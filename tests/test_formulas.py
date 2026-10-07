"""Formula edge cases.

Each case is checked twice: the reference evaluator must give the stated result, and a
workbook built with the same rows and recalculated by headless LibreOffice must give the
same value in every formula cell, with no error value anywhere."""

from __future__ import annotations

import copy
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from support import B, RM, W, needs_libreoffice, reset_root

AS_OF = "2026-10-05T23:59:59Z"


def task(n, status, due, completed_at=None, preds=(), start="2026-09-01"):
    t = {"task_id": f"TSK-{n:06d}", "project_id": "PRJ-000001", "phase_id": "PHS-000002", "milestone_id": "MLS-000002",
         "name": f"Edge case task {n}", "status": status, "owner_role_id": "ROL-000001", "planned_hours": 2,
         "predecessor_task_ids": list(preds), "planned_start_date": start, "critical_flag": False, "schema_version": "0.1.0"}
    if due:
        t["planned_due_date"] = due
    if completed_at:
        t["completed_at"] = completed_at
        t["actual_hours"] = 2
    return t


def risk(n, likelihood, impact):
    r = {"risk_id": f"RSK-{n:06d}", "project_id": "PRJ-000001", "phase_id": "PHS-000002", "title": f"Edge case risk {n}",
         "description": "Synthetic edge case.", "status": "open", "owner_role_id": "ROL-000001", "identified_at": "2026-10-01T10:00:00Z",
         "schema_version": "0.1.0"}
    if likelihood is not None:
        r["likelihood"] = likelihood
    if impact is not None:
        r["impact"] = impact
    return r


def request(n, status, changed, due=None):
    r = {"request_id": f"REQ-{n:06d}", "project_id": "PRJ-000001", "request_type": "information", "status": status,
         "submitted_at": "2026-10-01T10:00:00Z", "submitted_by_role_id": "ROL-000002", "owner_role_id": "ROL-000001",
         "priority": "normal", "status_changed_at": changed, "source": "customer", "description": "Synthetic edge case.",
         "schema_version": "0.1.0"}
    if due:
        r["escalation_due_at"] = due
    return r


def scorecard(rds, first_rde, met_fn, evidence=True, drop=None):
    rows = []
    totals = {"people": 4, "process": 4, "technology": 5, "data": 4, "training": 4, "support": 4}
    for i, (cat, total) in enumerate(totals.items()):
        if cat == drop:
            continue
        rows.append({"readiness_entry_id": f"RDE-{first_rde + i:06d}", "readiness_scorecard_id": rds, "project_id": "PRJ-000001",
                     "assessment_purpose": "pre_assessment", "assessed_at": "2026-10-02T10:00:00Z", "calculation_version": "1.0.0",
                     "category": cat, "criteria_total_count": total, "criteria_met_count": met_fn(total),
                     "required_evidence_complete_flag": evidence, "schema_version": "0.1.0"})
    return rows


def edge_inputs() -> dict:
    inputs = copy.deepcopy(W.load_inputs())
    inputs["tblTasks"] += [
        task(101, "not_started", None),                                     # blank due date
        task(102, "completed", "2026-10-03", "2026-10-02T12:00:00Z"),       # completed before due
        task(103, "completed", "2026-10-03", "2026-10-04T12:00:00Z"),       # completed after due
        task(104, "in_progress", "2026-10-09"),                             # open and before due
        task(105, "in_progress", "2026-10-01"),                             # open and past due
        task(106, "blocked", "2026-10-01"),                                 # blocked (even past due)
        task(107, "not_started", "2026-10-09", preds=["TSK-000101"]),       # valid predecessor
        task(108, "not_started", "2026-10-09", preds=["TSK-000999"]),       # invalid predecessor
        task(109, "not_started", "2026-10-05"),                             # due on the as-of day
    ]
    inputs["tblRisks"] += [risk(101, 1, 1), risk(102, 5, 5), risk(103, 6, 2), risk(104, 0, 3), risk(105, 2.5, 2), risk(106, None, 3)]
    inputs["tblRequests"] += [
        request(101, "in_progress", "2026-10-04T10:00:00Z"),                                   # no escalation timer
        request(102, "submitted", "2026-10-05T00:00:00Z", "2026-10-06T00:00:00Z"),             # due later
        request(103, "blocked", "2026-10-03T23:59:59Z", "2026-10-05T23:59:59Z"),               # due exactly at as-of
        request(104, "closed", "2026-10-04T10:00:00Z"),                                        # closed
        request(105, "awaiting_approval", "2026-10-01T00:00:00Z", "2026-10-04T01:00:00Z"),     # mismatch with the R1 rule
    ]
    inputs["tblReadiness"] += scorecard("RDS-000101", 101, lambda t: 0, evidence=False)        # score zero
    inputs["tblReadiness"] += scorecard("RDS-000102", 111, lambda t: t, evidence=True)         # full score
    inputs["tblReadiness"] += scorecard("RDS-000103", 121, lambda t: 1, drop="data")           # missing category
    return inputs


def by_id(ctx, table, row_id):
    spec = W.tables()[table]
    return next(r for r in ctx[table] if r[spec.id_column] == row_id)


class ReferenceEdgeCases(unittest.TestCase):
    """The reference evaluator states the intended result of every edge case."""

    @classmethod
    def setUpClass(cls) -> None:
        reset_root()
        cls.ctx = RM.compute(edge_inputs())

    def test_task_due_states(self) -> None:
        want = {101: "no_due_date", 102: "completed_on_time", 103: "completed_late", 104: "open_not_due",
                105: "overdue", 106: "blocked", 109: "open_not_due"}
        for n, state in want.items():
            self.assertEqual(by_id(self.ctx, "tblTasks", f"TSK-{n:06d}")["due_state"], state, n)

    def test_predecessors(self) -> None:
        self.assertEqual(by_id(self.ctx, "tblTasks", "TSK-000101")["predecessor_reference_valid"], "no_predecessor")
        self.assertEqual(by_id(self.ctx, "tblTasks", "TSK-000107")["predecessor_reference_valid"], "valid")
        self.assertEqual(by_id(self.ctx, "tblTasks", "TSK-000108")["predecessor_reference_valid"], "invalid")

    def test_risk_scale(self) -> None:
        low = by_id(self.ctx, "tblRisks", "RSK-000101")
        high = by_id(self.ctx, "tblRisks", "RSK-000102")
        self.assertEqual((low["risk_score"], low["risk_band"], low["risk_severity"]), (1, "low", "sev4"))
        self.assertEqual((high["risk_score"], high["risk_band"], high["risk_severity"]), (25, "critical", "sev1"))
        for n in (103, 104, 105, 106):
            r = by_id(self.ctx, "tblRisks", f"RSK-{n:06d}")
            self.assertEqual((r["risk_score"], r["risk_band"]), ("invalid_scale_value", "invalid_scale_value"), n)

    def test_request_escalation_states(self) -> None:
        want = {101: ("no_escalation_timer", "consistent"), 102: ("due_later", "consistent"),
                103: ("escalation_due", "consistent"), 104: ("closed", "consistent"), 105: ("escalation_due", "mismatch")}
        for n, (state, check) in want.items():
            r = by_id(self.ctx, "tblRequests", f"REQ-{n:06d}")
            self.assertEqual((r["escalation_due_state"], r["escalation_due_check"]), (state, check), n)

    def test_readiness_cases(self) -> None:
        summary = {r["readiness_scorecard_id"]: r for r in self.ctx["tblReadinessSummary"]}
        self.assertEqual(summary["RDS-000003"]["overall_score"], 91.25, "all categories present")
        self.assertEqual(summary["RDS-000101"]["overall_score"], 0, "score zero")
        self.assertIs(summary["RDS-000101"]["incomplete_required_evidence_flag"], True, "required evidence missing")
        self.assertEqual(summary["RDS-000102"]["overall_score"], 100, "full score")
        self.assertIs(summary["RDS-000102"]["incomplete_required_evidence_flag"], False)
        self.assertEqual(summary["RDS-000103"]["overall_score"], "incomplete_scorecard", "missing category")
        self.assertTrue(all(r["category_count_check"] == "missing_or_extra_category"
                            for r in self.ctx["tblReadiness"] if r["readiness_scorecard_id"] == "RDS-000103"))

    def test_weights_valid(self) -> None:
        weights = self.ctx["tblReadinessWeights"]
        total = next(p["value"] for p in self.ctx["tblRuleParameters"] if p["parameter"] == "readiness_total_weight")
        self.assertEqual(sum(w["weight"] for w in weights), total)
        self.assertEqual([w["category"] for w in weights], ["people", "process", "technology", "data", "training", "support"])


class EmptyWorkbookReference(unittest.TestCase):
    def test_empty_tables_give_zeros_not_errors(self) -> None:
        reset_root()
        ctx = RM.compute({k: [] for k in W.load_inputs()})
        kpi = {r["kpi_key"]: r["value"] for r in ctx["tblKpiSummary"]}
        self.assertTrue(all(v == 0 for v in kpi.values()), kpi)
        cap = {r["check_key"]: r["value"] for r in ctx["tblCapacitySummary"]}
        self.assertEqual(cap["hours_reconciliation"], "reconciled")
        self.assertTrue(all(r["readiness_score"] == "" for r in ctx["tblReadinessByProject"]))


def recalc_compare(case: unittest.TestCase, inputs: dict) -> None:
    reset_root()
    with tempfile.TemporaryDirectory() as tmp:
        xlsx = Path(tmp) / "case.xlsx"
        B.build(xlsx, inputs)
        recalculated = B.recalculate(xlsx, Path(tmp) / "out")
        wb = load_workbook(recalculated, data_only=True)
        errors = [f"{ws.title}!{c.coordinate}={c.value}" for ws in wb.worksheets for row in ws.iter_rows() for c in row
                  if isinstance(c.value, str) and c.value.startswith(("#", "Err:"))]
        case.assertEqual(errors, [], "error values after recalculation")
        got = B.read_tables(recalculated)
        ctx = RM.compute(inputs)
        import verify_workbook as V

        mismatches = []
        for name, spec in W.tables().items():
            for column in spec.columns:
                if column.kind != "formula":
                    continue
                expected_rows = ctx[name] or [{}]
                for i, (g, e) in enumerate(zip(got[name], expected_rows)):
                    if not V.same(g.get(column.name), e.get(column.name, "")):
                        mismatches.append(f"{name}[{i + 1}].{column.name}: workbook {g.get(column.name)!r}, reference {e.get(column.name)!r}")
        case.assertEqual(mismatches, [])


class RecalculatedEdgeCases(unittest.TestCase):
    @needs_libreoffice
    def test_edge_case_workbook_matches_reference(self) -> None:
        recalc_compare(self, edge_inputs())

    @needs_libreoffice
    def test_empty_workbook_has_no_error_values(self) -> None:
        recalc_compare(self, {k: [] for k in W.load_inputs()})


if __name__ == "__main__":
    unittest.main()
