"""Tests for the builder, verifier, and pin synchronization tool."""

from __future__ import annotations

import filecmp
import tempfile
import unittest
from pathlib import Path

from support import B, ROOT, Sandbox, W, reset_root  # isort: skip  (must precede openpyxl)

from openpyxl import load_workbook  # noqa: E402

import sync_r1_contracts as S  # noqa: E402
import verify_workbook as V  # noqa: E402

R1_CLONE = ROOT.parent / "implementation-operating-system"


class Builder(unittest.TestCase):
    def setUp(self) -> None:
        reset_root()

    def test_build_is_byte_identical_and_equals_committed_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            a, b = Path(tmp) / "a.xlsx", Path(tmp) / "b.xlsx"
            B.build(a)
            B.build(b)
            self.assertTrue(filecmp.cmp(a, b, shallow=False), "two builds differ")
            self.assertTrue(filecmp.cmp(a, W.WORKBOOK_PATH, shallow=False), "the committed workbook is not the build output")

    def test_build_does_not_depend_on_optional_lxml(self) -> None:
        # Regression: CI (no lxml) produced different bytes than a machine with lxml,
        # because openpyxl switches XML writers when lxml is installed.
        import openpyxl.xml

        self.assertFalse(openpyxl.xml.LXML, "the standard XML writer must be in effect")
        original = B.openpyxl_xml.LXML
        try:
            B.openpyxl_xml.LXML = True
            with tempfile.TemporaryDirectory() as tmp, self.assertRaises(RuntimeError):
                B.build(Path(tmp) / "x.xlsx")
        finally:
            B.openpyxl_xml.LXML = original

    def test_twelve_visible_sheets_in_order(self) -> None:
        wb = load_workbook(W.WORKBOOK_PATH)
        self.assertEqual(wb.sheetnames, list(W.SHEETS))
        self.assertTrue(all(ws.sheet_state == "visible" for ws in wb.worksheets))

    def test_named_settings_and_lists_exist(self) -> None:
        wb = load_workbook(W.WORKBOOK_PATH)
        for name in ("calculation_as_of_at", "event_reconciliation_from_at", "selected_org_stage", "operating_mode",
                     "risk_likelihood_min", "risk_likelihood_max", "risk_impact_min", "risk_impact_max", "lst_request_status"):
            self.assertIn(name, wb.defined_names)

    def test_readme_tab_carries_labels_and_statement(self) -> None:
        wb = load_workbook(W.WORKBOOK_PATH)
        text = " ".join(str(c.value) for ws in (wb["README"], wb["Readiness Scorecard"]) for row in ws.iter_rows() for c in row if c.value)
        for phrase in ("synthetic data", "illustrative example", "user-configurable parameter",
                       "Readiness score informs human review. It does not decide go or no-go."):
            self.assertIn(phrase, text)
        self.assertIn("Source of truth: R1 configuration.", str(wb["Escalation Rules"]["A2"].value))

    def test_document_metadata(self) -> None:
        wb = load_workbook(W.WORKBOOK_PATH)
        self.assertEqual(wb.properties.creator, "Jared Muldrow")
        self.assertEqual(wb.properties.lastModifiedBy, "Jared Muldrow")

    def test_structural_verification_passes(self) -> None:
        report = V.Report()
        V.structural_checks(W.WORKBOOK_PATH, report)
        self.assertEqual([r for r in report.rows if r[1] != "PASS"], [])


class Sync(unittest.TestCase):
    def test_refuses_branch_names_and_latest(self) -> None:
        for bad in ("main", "latest", "HEAD", "9acd25a"):
            with self.assertRaises(SystemExit):
                S.main(["--r1-repo", str(ROOT), "--r1-commit", bad])

    @unittest.skipUnless((R1_CLONE / ".git").exists(), "a local R1 clone is needed to re-run the sync")
    def test_resync_from_the_pinned_commit_changes_nothing(self) -> None:
        box = Sandbox()
        original = S.ROOT
        try:
            S.ROOT = box.root
            commit = W.pin_reference()["r1_commit"]
            S.main(["--r1-repo", str(R1_CLONE), "--r1-commit", commit])
            for rel in ("standard/standard-reference.yaml", "schemas/r1/manifest.yaml", "data/synthetic/manifest.yaml"):
                self.assertEqual((box.root / rel).read_bytes(), (ROOT / rel).read_bytes(), rel)
        finally:
            S.ROOT = original
            box.close()


if __name__ == "__main__":
    unittest.main()
